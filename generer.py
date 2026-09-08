import openpyxl
import json

wb = openpyxl.load_workbook("points.xlsx")
wb_2 = openpyxl.load_workbook("numeros_eleves.xlsx")

COULEURS = {
    "Pingouins_bresiliens": "#1b581b",
    "Les_11_physiciens": "#173253",
    "Panthere_rose": "#bf328f",
    "Miel_pops": "#e4892a",
    "Griffondor": "#a9254d",
    "Team_3": "#1800ad",
    "La_vague": "#335b9c",
    "Cerveaux_lents": "#e9a6a3",
    "Bouzelouf": "#a96127",
}

resultat = {}

for nom_classe in wb.sheetnames:
    feuille = wb[nom_classe]
    feuille_2 = wb_2[nom_classe]
    totaux_maisons = {}
    totaux_eleves = {}

    for ligne in feuille.iter_rows(min_row = 2, values_only=True): 
        nom, prenom, maison, points = ligne
        
        if nom is None or prenom is None :
            continue
        nom = nom.strip()
        prenom = prenom.strip()
        maison = maison.strip()

        try :
            points= int(points)
        except : 
            points = 0

        eleve = f"{prenom.capitalize()} {nom.upper()}"    
        totaux_maisons[maison] = totaux_maisons.get(maison, 0) + points
        
        for ligne_2 in feuille_2.iter_rows(min_row = 2, values_only=True):
            nom_2, prenom_2, numero = ligne_2
            
            if nom_2 is None :
                continue

            eleve_2 = f"{prenom_2.capitalize()} {nom_2.upper()}"

            numero_eleve = None
            if eleve == eleve_2 :
                numero_eleve = numero
                break

        if eleve not in totaux_eleves:
            totaux_eleves[eleve] = {"points": 0, "maison": maison, "numéro": numero_eleve}
        totaux_eleves[eleve]["points"] = totaux_eleves[eleve]["points"] + points
    

    liste_maisons = []
    for nom_maison, total in totaux_maisons.items():
        liste_maisons.append({
            "nom": nom_maison,
            "couleur": COULEURS.get(nom_maison,"#888888"),  # gris si maison inconnue
            "total": total
        })

    liste_eleves = []
    for nom_eleve, infos in totaux_eleves.items():
        liste_eleves.append({
            "numéro": infos["numéro"],
            "points": infos["points"],
            "maison": infos["maison"],
            "couleur": COULEURS.get(infos["maison"], "#888888")
        })

    resultat[nom_classe] = {
    "maisons": liste_maisons,
    "eleves": liste_eleves
}

with open("data.json", "w", encoding="utf-8") as f:
    json.dump(resultat, f, ensure_ascii=False, indent=2)
